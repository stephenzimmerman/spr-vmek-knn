#!/usr/bin/env python3
"""
SPR VMEK KNN Processor — Standalone GUI v2

Changes in v2:
- Folder batch mode creates ONE output CSV only.
- No individual per-input output files are created.
- Output includes SampleName from the input file name.
- Output includes Barcode extracted as 11 characters beginning with LC or DR.
- Suppresses sklearn feature-name warnings during prediction.

Required packages:
  pip install pandas numpy scikit-learn joblib openpyxl
"""

from __future__ import annotations

import json
import math
import os
import queue
import re
import subprocess
import sys
import threading
import traceback
import warnings
from pathlib import Path
from typing import Any, Dict, Tuple

import joblib
import numpy as np
import pandas as pd
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

DEFAULT_MACHINE_REGISTRY: Dict[str, Dict[str, Any]] = {
    "Quality Lab283": {"machine_name": "QL2020283", "site_apl": "Quality Lab283", "country": "USA", "vmek_model": "MTX–170–36", "camera": "Basler acA1300-60gc", "px_mm": 0.332, "area_corr": 0.93513, "bias_flag": "— Monitor"},
    "Quality Lab241": {"machine_name": "QL2019241", "site_apl": "Quality Lab241", "country": "USA", "vmek_model": "MTX–170–36", "camera": "Basler acA1300-60gc", "px_mm": 0.332, "area_corr": 0.99476, "bias_flag": "✓ Good"},
    "CLGR APL2": {"machine_name": "APLCLGR02AL23", "site_apl": "CLGR APL2", "country": "Chile", "vmek_model": "MX3–170–36", "camera": "JAI GOX-2402C-PGE", "px_mm": 0.279, "area_corr": 1.02843, "bias_flag": "✓ Good"},
    "MXPV APL1": {"machine_name": "APLMXPV01AL19", "site_apl": "MXPV APL1", "country": "Mexico", "vmek_model": "MX3–170–36", "camera": "JAI GOX-2402C-PGE", "px_mm": 0.279, "area_corr": 1.14804, "bias_flag": "⚠ Review"},
    "MXPV APL2": {"machine_name": "APLMXPV02AL21", "site_apl": "MXPV APL2", "country": "Mexico", "vmek_model": "MTX–170–36", "camera": "Basler acA1300-60gc", "px_mm": 0.332, "area_corr": 1.08827, "bias_flag": "⚠ Review"},
    "MXPV APL3": {"machine_name": "APLMXPV03AL23", "site_apl": "MXPV APL3", "country": "Mexico", "vmek_model": "MX3–170–36", "camera": "JAI GOX-2402C-PGE", "px_mm": 0.279, "area_corr": 1.06068, "bias_flag": "— Monitor"},
    "PRSA APL1": {"machine_name": "APLPRSA01AL19", "site_apl": "PRSA APL1", "country": "USA", "vmek_model": "MTX–170–36", "camera": "Basler acA1300-60gc", "px_mm": 0.332, "area_corr": 0.90205, "bias_flag": "⚠ Review"},
    "PRSA APL2": {"machine_name": "APLPRSA02AL23", "site_apl": "PRSA APL2", "country": "USA", "vmek_model": "MX3–170–36", "camera": "JAI GOX-2402C-PGE", "px_mm": 0.279, "area_corr": 1.08019, "bias_flag": "— Monitor"},
    "USSL APL1": {"machine_name": "APLUSSL01AL17", "site_apl": "USSL APL1", "country": "USA", "vmek_model": "MTX–170–36", "camera": "Basler acA1300-60gc", "px_mm": 0.332, "area_corr": 0.91864, "bias_flag": "⚠ Review"},
    "USSL APL2": {"machine_name": "APLUSSL02AL19", "site_apl": "USSL APL2", "country": "USA", "vmek_model": "MTX–170–36", "camera": "Basler acA1300-60gc", "px_mm": 0.332, "area_corr": 0.94023, "bias_flag": "— Monitor"},
    "Elk Farm SPR": {"machine_name": "Elk Farm", "site_apl": "Elk Farm SPR", "country": "USA", "vmek_model": "MTX–170–36", "camera": "Basler acA1300-60gc", "px_mm": 0.332, "area_corr": 1.08911, "bias_flag": "⚠ Review"},
}

FEATURES = ["Area_mm2_corrected", "Length_mm_corrected", "Width_mm_corrected", "ContLength_mm_equiv_corrected"]
BUCKET_ORDER = ["IPSDP", "IPSSP", "IPLSP", "IPLDP"]
SUPPORTED_SUFFIXES = {".csv", ".txt", ".tsv", ".xlsx", ".xlsm", ".xls"}
LOCAL_REGISTRY_FILE = "machine_registry.json"
DEFAULT_MODEL_NAME = "vmek_knn_final_optionB_plus_2021_IPLDP.joblib"
DEFAULT_OUTPUT_NAME = "vmek_spr_traits_output.csv"

# Hide noisy sklearn feature-name warnings in local GUI runs.
warnings.filterwarnings("ignore", message=".*valid feature names.*", category=UserWarning)
warnings.filterwarnings("ignore", message=".*feature names.*", category=UserWarning)


class VmekError(Exception):
    pass


def validate_registry(registry: Dict[str, Dict[str, Any]]) -> None:
    for site, rec in registry.items():
        if not site or not rec.get("machine_name"):
            raise VmekError(f"Invalid registry entry: {site}")
        rec["px_mm"] = float(rec["px_mm"])
        rec["area_corr"] = float(rec["area_corr"])
        if rec["px_mm"] <= 0 or rec["area_corr"] <= 0:
            raise VmekError(f"Invalid correction values for {site}")


def load_registry(path: str | Path = LOCAL_REGISTRY_FILE) -> Dict[str, Dict[str, Any]]:
    p = Path(path)
    registry = json.loads(p.read_text(encoding="utf-8")) if p.exists() else json.loads(json.dumps(DEFAULT_MACHINE_REGISTRY))
    validate_registry(registry)
    return registry


def save_registry(registry: Dict[str, Dict[str, Any]], path: str | Path = LOCAL_REGISTRY_FILE) -> None:
    validate_registry(registry)
    Path(path).write_text(json.dumps(registry, indent=2, ensure_ascii=False), encoding="utf-8")


def norm_header(x: Any) -> str:
    return re.sub(r"\s+", " ", str(x).replace("\n", " ")).strip()


def update_registry_from_excel(excel_path: str | Path, registry_path: str | Path = LOCAL_REGISTRY_FILE) -> Dict[str, Dict[str, Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(excel_path, data_only=True)
    if "Standardization Table" not in wb.sheetnames:
        raise VmekError("Workbook must contain a 'Standardization Table' sheet.")
    ws = wb["Standardization Table"]
    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(4, c).value
        if v:
            headers[norm_header(v)] = c
    required = ["Machine Name", "Site / APL", "px→mm Factor", "Correction Factor"]
    missing = [h for h in required if h not in headers]
    if missing:
        raise VmekError(f"Missing required columns in standardization table: {missing}")
    registry = {}
    for r in range(5, ws.max_row + 1):
        machine = ws.cell(r, headers["Machine Name"]).value
        site = ws.cell(r, headers["Site / APL"]).value
        px = ws.cell(r, headers["px→mm Factor"]).value
        corr = ws.cell(r, headers["Correction Factor"]).value
        if not machine or not site or px in (None, "") or corr in (None, ""):
            continue
        site = str(site).strip()
        if site in registry:
            raise VmekError(f"Duplicate Site / APL in workbook: {site}")
        registry[site] = {
            "machine_name": str(machine).strip(),
            "site_apl": site,
            "country": ws.cell(r, headers.get("Country", 0)).value if "Country" in headers else None,
            "vmek_model": ws.cell(r, headers.get("VMEK Model", 0)).value if "VMEK Model" in headers else None,
            "camera": ws.cell(r, headers.get("Camera", 0)).value if "Camera" in headers else None,
            "px_mm": float(px),
            "area_corr": float(corr),
            "bias_flag": ws.cell(r, headers.get("Bias Flag", 0)).value if "Bias Flag" in headers else None,
        }
    save_registry(registry, registry_path)
    return registry


def extract_barcode_from_filename(path: str | Path) -> str:
    """Extract 11-char barcode starting with LC or DR from input file name/stem."""
    stem = Path(path).stem.upper()
    match = re.search(r"(?:LC|DR)[A-Z0-9]{9}", stem)
    return match.group(0) if match else ""


def sample_name_from_filename(path: str | Path) -> str:
    return Path(path).stem


def read_vmek_file(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(path)
    return pd.read_csv(path, sep=None, engine="python")


def detect_schema(df: pd.DataFrame) -> str:
    cols = set(df.columns)
    if {"Area", "Length", "Width", "ContLength"}.issubset(cols):
        return "legacy_px"
    if {"Area", "Length", "Width", "Roundness"}.issubset(cols):
        return "mx4_mm"
    raise VmekError("Could not detect schema. Need Area/Length/Width/ContLength or Area/Length/Width/Roundness.")


def preprocess_option_b(df: pd.DataFrame, site_apl: str, registry: Dict[str, Dict[str, Any]], legacy_filter: bool = True) -> Tuple[pd.DataFrame, Dict[str, Any]]:
    if site_apl not in registry:
        raise VmekError(f"Site/APL not found in registry: {site_apl}")
    raw = df.copy()
    total_raw = len(raw)
    schema = detect_schema(raw)
    rec = registry[site_apl]
    px_mm = float(rec["px_mm"])
    area_corr = float(rec["area_corr"])
    linear_corr = math.sqrt(area_corr)

    if "Id" not in raw.columns:
        raw["Id"] = np.arange(1, len(raw) + 1)
    for c in ["Area", "Length", "Width", "ContLength", "Roundness"]:
        if c in raw.columns:
            raw[c] = pd.to_numeric(raw[c], errors="coerce")

    required = ["Area", "Length", "Width"] + (["ContLength"] if schema == "legacy_px" else ["Roundness"])
    valid = raw.dropna(subset=required).copy()

    removed_area = 0
    removed_roundness = 0
    if schema == "legacy_px" and legacy_filter:
        before = len(valid)
        valid = valid[(valid["Area"] >= 100) & (valid["Area"] <= 1000)].copy()
        removed_area = before - len(valid)
    if schema == "mx4_mm":
        before = len(valid)
        valid = valid[valid["Roundness"] > 0].copy()
        removed_roundness = before - len(valid)

    if schema == "legacy_px":
        valid["Area_mm2_corrected"] = valid["Area"] * (px_mm ** 2) * area_corr
        valid["Length_mm_corrected"] = valid["Length"] * px_mm * linear_corr
        valid["Width_mm_corrected"] = valid["Width"] * px_mm * linear_corr
        valid["ContLength_mm_equiv_corrected"] = valid["ContLength"] * px_mm * linear_corr
    else:
        valid["Area_mm2_corrected"] = valid["Area"] * area_corr
        valid["Length_mm_corrected"] = valid["Length"] * linear_corr
        valid["Width_mm_corrected"] = valid["Width"] * linear_corr
        valid["ContLength_mm_equiv_corrected"] = np.sqrt((4 * math.pi * valid["Area_mm2_corrected"]) / valid["Roundness"])

    stats = {
        "total_raw": int(total_raw),
        "valid_kernels": int(len(valid)),
        "kernels_removed": int(total_raw - len(valid)),
        "removed_area_filter": int(removed_area),
        "removed_bad_roundness": int(removed_roundness),
        "schema_detected": schema,
        "site_apl": site_apl,
        "machine_name": rec.get("machine_name"),
    }
    return valid, stats


def load_model(path: str | Path):
    pkg = joblib.load(path)
    if isinstance(pkg, dict) and "pipeline" in pkg:
        return pkg["pipeline"], pkg.get("metadata", {})
    return pkg, {}


def summarize_one_file(input_path: str | Path, site_apl: str, model, metadata: Dict[str, Any], registry: Dict[str, Dict[str, Any]], legacy_filter: bool = True) -> Dict[str, Any]:
    sample_name = sample_name_from_filename(input_path)
    barcode = extract_barcode_from_filename(input_path)
    base = {
        "SampleName": sample_name,
        "Barcode": barcode,
        "InputFile": Path(input_path).name,
        "Site_APL": site_apl,
        "Machine_Name": registry.get(site_apl, {}).get("machine_name", ""),
        "schema_detected": "",
        "iSCLCN": 0,
        "total_kernels_raw_file": 0,
        "kernels_removed_file": 0,
        "IPSDP": np.nan,
        "IPSSP": np.nan,
        "IPLSP": np.nan,
        "IPLDP": np.nan,
        "Processing_Status": "failed",
        "Error_Message": "",
    }
    try:
        raw = read_vmek_file(input_path)
        kernels, stats = preprocess_option_b(raw, site_apl, registry, legacy_filter)
        if kernels.empty:
            raise VmekError("No valid kernels remained after preprocessing.")
        features = metadata.get("features", FEATURES)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            preds = model.predict(kernels[features])
        counts = pd.Series(preds).value_counts().reindex(BUCKET_ORDER, fill_value=0)
        total = int(counts.sum())
        pct = (counts / total * 100).round(3) if total else counts.astype(float)
        base.update({
            "schema_detected": stats["schema_detected"],
            "iSCLCN": total,
            "total_kernels_raw_file": stats["total_raw"],
            "kernels_removed_file": stats["kernels_removed"],
            "IPSDP": float(pct["IPSDP"]),
            "IPSSP": float(pct["IPSSP"]),
            "IPLSP": float(pct["IPLSP"]),
            "IPLDP": float(pct["IPLDP"]),
            "Processing_Status": "success",
            "Error_Message": "",
        })
    except Exception as exc:
        base["Error_Message"] = str(exc)
    return base


def find_files(folder: str | Path, recursive: bool = False) -> list[Path]:
    folder = Path(folder)
    iterator = folder.rglob("*") if recursive else folder.glob("*")
    skip = ["_spr_traits", "_kernel_predictions", "_processing_stats", "batch_processing", "vmek_spr_traits_output"]
    return sorted([p for p in iterator if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES and not any(s in p.stem for s in skip)])


def predict_folder_single_output(input_folder: str | Path, site_apl: str, model_path: str | Path, output_dir: str | Path, output_name: str, registry_path: str | Path = LOCAL_REGISTRY_FILE, recursive: bool = False, stop_on_error: bool = False, legacy_filter: bool = True) -> Path:
    registry = load_registry(registry_path)
    files = find_files(input_folder, recursive)
    if not files:
        raise VmekError(f"No supported files found in folder: {input_folder}")
    if not output_name.lower().endswith(".csv"):
        output_name += ".csv"
    model, metadata = load_model(model_path)
    rows = []
    for f in files:
        row = summarize_one_file(f, site_apl, model, metadata, registry, legacy_filter)
        rows.append(row)
        if row["Processing_Status"] == "failed" and stop_on_error:
            break
    outdir = Path(output_dir)
    outdir.mkdir(parents=True, exist_ok=True)
    outpath = outdir / output_name
    cols = ["SampleName", "Barcode", "InputFile", "Site_APL", "Machine_Name", "schema_detected", "iSCLCN", "total_kernels_raw_file", "kernels_removed_file", "IPSDP", "IPSSP", "IPLSP", "IPLDP", "Processing_Status", "Error_Message"]
    pd.DataFrame(rows)[cols].to_csv(outpath, index=False)
    return outpath


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("SPR VMEK KNN Processor")
        self.geometry("980x640")
        self.registry_path = tk.StringVar(value=LOCAL_REGISTRY_FILE)
        here = Path(__file__).resolve().parent
        self.model_path = tk.StringVar(value=str(here / DEFAULT_MODEL_NAME))
        self.input_folder = tk.StringVar(value="")
        self.output_folder = tk.StringVar(value=str(here / "vmek_outputs"))
        self.site_apl = tk.StringVar(value="")
        self.recursive = tk.BooleanVar(value=False)
        self.stop_on_error = tk.BooleanVar(value=False)
        self.no_filter = tk.BooleanVar(value=False)
        self.output_name = tk.StringVar(value=DEFAULT_OUTPUT_NAME)
        self.q = queue.Queue()
        self.worker = None
        self._ui()
        self.refresh_machines()
        self.after(200, self.poll)

    def _ui(self):
        root = ttk.Frame(self, padding=12)
        root.pack(fill="both", expand=True)
        root.columnconfigure(0, weight=1)
        root.rowconfigure(4, weight=1)
        ttk.Label(root, text="SPR VMEK KNN Processor", font=("Segoe UI", 16, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(root, text="Folder batch input → one data-system CSV output", foreground="#444").grid(row=0, column=0, sticky="e")

        mf = ttk.LabelFrame(root, text="1. Machine", padding=10)
        mf.grid(row=1, column=0, sticky="ew", pady=6)
        mf.columnconfigure(1, weight=1)
        ttk.Label(mf, text="Site / APL").grid(row=0, column=0, sticky="w")
        self.combo = ttk.Combobox(mf, textvariable=self.site_apl, state="readonly", width=35)
        self.combo.grid(row=0, column=1, sticky="w", padx=8)
        self.combo.bind("<<ComboboxSelected>>", lambda e: self.show_machine())
        ttk.Button(mf, text="Refresh", command=self.refresh_machines).grid(row=0, column=2, padx=5)
        ttk.Button(mf, text="Update from Excel...", command=self.update_excel).grid(row=0, column=3, padx=5)
        self.machine_info = ttk.Label(mf, text="")
        self.machine_info.grid(row=1, column=0, columnspan=4, sticky="w", pady=(8, 0))

        io = ttk.LabelFrame(root, text="2. Folder batch input", padding=10)
        io.grid(row=2, column=0, sticky="ew", pady=6)
        io.columnconfigure(1, weight=1)
        self.path_row(io, 0, "Raw data folder", self.input_folder, lambda: self.pick_dir(self.input_folder))
        self.path_row(io, 1, "Output folder", self.output_folder, lambda: self.pick_dir(self.output_folder))
        self.path_row(io, 2, "Model .joblib", self.model_path, self.pick_model)
        ttk.Label(io, text="Single output CSV name").grid(row=3, column=0, sticky="w")
        ttk.Entry(io, textvariable=self.output_name).grid(row=3, column=1, sticky="ew", padx=8, pady=4)

        opt = ttk.LabelFrame(root, text="3. Options", padding=10)
        opt.grid(row=3, column=0, sticky="ew", pady=6)
        ttk.Checkbutton(opt, text="Search subfolders", variable=self.recursive).grid(row=0, column=0, sticky="w", padx=5)
        ttk.Checkbutton(opt, text="Stop on first error", variable=self.stop_on_error).grid(row=0, column=1, sticky="w", padx=5)
        ttk.Checkbutton(opt, text="Disable legacy 100–1000 px filter", variable=self.no_filter).grid(row=0, column=2, sticky="w", padx=5)

        run = ttk.LabelFrame(root, text="4. Run", padding=10)
        run.grid(row=4, column=0, sticky="nsew", pady=6)
        run.columnconfigure(0, weight=1)
        run.rowconfigure(1, weight=1)
        bar = ttk.Frame(run)
        bar.grid(row=0, column=0, sticky="ew")
        self.run_btn = ttk.Button(bar, text="Process folder", command=self.start)
        self.run_btn.pack(side="left")
        ttk.Button(bar, text="Open output folder", command=self.open_output).pack(side="left", padx=8)
        ttk.Button(bar, text="Clear log", command=lambda: self.log.delete("1.0", "end")).pack(side="left")
        self.progress = ttk.Progressbar(bar, mode="indeterminate", length=220)
        self.progress.pack(side="right")
        self.log = tk.Text(run, height=16, wrap="word")
        self.log.grid(row=1, column=0, sticky="nsew", pady=(8, 0))
        scroll = ttk.Scrollbar(run, command=self.log.yview)
        scroll.grid(row=1, column=1, sticky="ns", pady=(8, 0))
        self.log.configure(yscrollcommand=scroll.set)
        self.write("Ready. Select folder, machine, model, then Process folder.")
        self.write("Output is one CSV only. Barcode is extracted from file names like Sample_DR287636625...")

    def path_row(self, parent, row, label, var, cmd):
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w")
        ttk.Entry(parent, textvariable=var).grid(row=row, column=1, sticky="ew", padx=8, pady=4)
        ttk.Button(parent, text="Browse...", command=cmd).grid(row=row, column=2, pady=4)

    def write(self, msg):
        self.log.insert("end", str(msg).rstrip() + "\n")
        self.log.see("end")

    def refresh_machines(self):
        try:
            self.registry = load_registry(self.registry_path.get())
            sites = sorted(self.registry)
            self.combo["values"] = sites
            if self.site_apl.get() not in sites:
                self.site_apl.set(sites[0] if sites else "")
            self.show_machine()
            self.write(f"Loaded {len(sites)} machines.")
        except Exception as exc:
            messagebox.showerror("Registry error", str(exc))

    def show_machine(self):
        rec = getattr(self, "registry", {}).get(self.site_apl.get(), {})
        self.machine_info.configure(text=f"Machine: {rec.get('machine_name','')} | px→mm: {rec.get('px_mm','')} | area_corr: {rec.get('area_corr','')} | bias: {rec.get('bias_flag','')}")

    def update_excel(self):
        path = filedialog.askopenfilename(title="Select standardization workbook", filetypes=[("Excel", "*.xlsx *.xlsm *.xls"), ("All", "*.*")])
        if not path:
            return
        try:
            reg = update_registry_from_excel(path, self.registry_path.get())
            self.write(f"Updated registry from {path}; machines: {len(reg)}")
            self.refresh_machines()
        except Exception as exc:
            messagebox.showerror("Update failed", str(exc))

    def pick_dir(self, var):
        path = filedialog.askdirectory()
        if path:
            var.set(path)

    def pick_model(self):
        path = filedialog.askopenfilename(filetypes=[("Joblib", "*.joblib"), ("All", "*.*")])
        if path:
            self.model_path.set(path)

    def start(self):
        if self.worker and self.worker.is_alive():
            return
        if not Path(self.input_folder.get()).is_dir():
            messagebox.showerror("Missing input", "Select a valid raw data folder.")
            return
        if not Path(self.model_path.get()).exists():
            messagebox.showerror("Missing model", "Select the final .joblib model file.")
            return
        self.run_btn.configure(state="disabled")
        self.progress.start(10)
        self.write("\nStarting batch...")
        args = dict(
            input_folder=self.input_folder.get(),
            site_apl=self.site_apl.get(),
            model_path=self.model_path.get(),
            output_dir=self.output_folder.get(),
            output_name=self.output_name.get().strip() or DEFAULT_OUTPUT_NAME,
            registry_path=self.registry_path.get(),
            recursive=self.recursive.get(),
            stop_on_error=self.stop_on_error.get(),
            legacy_filter=not self.no_filter.get(),
        )
        self.worker = threading.Thread(target=self.run_worker, args=(args,), daemon=True)
        self.worker.start()

    def run_worker(self, args):
        try:
            outpath = predict_folder_single_output(**args)
            self.q.put(("ok", outpath))
        except Exception as exc:
            self.q.put(("err", (str(exc), traceback.format_exc())))

    def poll(self):
        try:
            while True:
                status, payload = self.q.get_nowait()
                self.progress.stop()
                self.run_btn.configure(state="normal")
                if status == "ok":
                    self.write("Batch complete.")
                    self.write(f"Output CSV: {payload}")
                    messagebox.showinfo("Complete", f"Batch processing complete.\n\nOutput:\n{payload}")
                else:
                    msg, tb = payload
                    self.write("ERROR: " + msg)
                    self.write(tb)
                    messagebox.showerror("Error", msg)
        except queue.Empty:
            pass
        self.after(200, self.poll)

    def open_output(self):
        p = Path(self.output_folder.get())
        p.mkdir(parents=True, exist_ok=True)
        if sys.platform.startswith("win"):
            os.startfile(p)  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.run(["open", str(p)], check=False)
        else:
            subprocess.run(["xdg-open", str(p)], check=False)


def main():
    App().mainloop()


if __name__ == "__main__":
    main()
